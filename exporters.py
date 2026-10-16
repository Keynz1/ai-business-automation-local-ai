from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook, load_workbook

from models import ProcessedRequest

HEADERS = [
    "request_id", "received_at", "customer_name", "company", "email",
    "request_type", "priority", "summary", "requested_action", "confidence",
    "api_status", "response_text", "processing_status",
]


def append_csv(items: list[ProcessedRequest], path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    exists = path.exists() and path.stat().st_size > 0
    with path.open("a", newline="", encoding="utf-8-sig") as handle:
        writer = csv.DictWriter(handle, fieldnames=HEADERS, extrasaction="ignore")
        if not exists:
            writer.writeheader()
        for item in items:
            writer.writerow(item.to_row())
    return path


def append_excel(items: list[ProcessedRequest], path: Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook.active
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Requests"
        sheet.append(HEADERS)

    for item in items:
        row = item.to_row()
        sheet.append([row.get(name, "") for name in HEADERS])

    workbook.save(path)
    return path
