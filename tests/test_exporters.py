from pathlib import Path

from openpyxl import load_workbook

from exporters import append_csv, append_excel
from models import ProcessedRequest, normalize_payload


def make_item(email=""):
    base = normalize_payload({"customer_name": "Ada", "email": email, "summary": "Need help"})
    return ProcessedRequest.from_business_request(
        base,
        api_status="demo-local",
        response_text="Thanks",
        processing_status="processed",
    )


def test_csv_creates_file_with_stable_header_and_empty_email(tmp_path: Path):
    path = append_csv([make_item("")], tmp_path / "out" / "requests.csv")
    text = path.read_text(encoding="utf-8-sig")
    lines = text.splitlines()
    assert lines[0].startswith("request_id,received_at,customer_name,company,email,request_type")
    assert len(lines) == 2


def test_csv_second_run_appends_without_duplicate_header(tmp_path: Path):
    path = tmp_path / "requests.csv"
    append_csv([make_item("a@example.com")], path)
    append_csv([make_item("b@example.com")], path)
    lines = path.read_text(encoding="utf-8-sig").splitlines()
    assert len(lines) == 3
    assert sum(line.startswith("request_id,") for line in lines) == 1


def test_excel_creates_and_appends_rows(tmp_path: Path):
    path = tmp_path / "requests.xlsx"
    append_excel([make_item("")], path)
    append_excel([make_item("b@example.com")], path)
    wb = load_workbook(path)
    ws = wb.active
    assert ws.max_row == 3
    assert ws.cell(1, 1).value == "request_id"
    assert ws.cell(2, 5).value in (None, "")
    assert ws.cell(3, 5).value == "b@example.com"
