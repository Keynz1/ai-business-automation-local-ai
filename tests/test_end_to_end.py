import csv
from pathlib import Path

from openpyxl import load_workbook

import processor
import run_demo
from models import normalize_payload


def deterministic_classify(text, model="qwen2:1.5b"):
    lower = text.lower()
    if "clean" in lower or "csv" in lower:
        kind = "data_cleanup"
    elif "lead" in lower or "api" in lower:
        kind = "automation"
    else:
        kind = "reporting"
    return normalize_payload({
        "customer_name": "Demo Client",
        "company": "Example Co",
        "email": "demo@example.com",
        "request_type": kind,
        "priority": "urgent" if "urgent" in lower else "normal",
        "summary": text,
        "requested_action": "Prepare solution",
        "confidence": 0.95,
    })


def test_three_request_demo_creates_all_outputs(monkeypatch, tmp_path: Path):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(processor, "classify_request", deterministic_classify)
    monkeypatch.setattr(processor, "enrich_request", lambda req, endpoint=None: {"api_status": "demo-local", "enrichment_note": "Local demo"})

    assert run_demo.main() == 0

    csv_path = tmp_path / "output" / "requests.csv"
    xlsx_path = tmp_path / "output" / "requests.xlsx"
    report_path = tmp_path / "output" / "processing_report.txt"
    assert csv_path.exists() and xlsx_path.exists() and report_path.exists()

    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 3
    assert all(row["processing_status"] == "processed" for row in rows)

    wb = load_workbook(xlsx_path)
    assert wb.active.max_row == 4
    assert report_path.read_text(encoding="utf-8").count("Request ID:") == 3
